"""Page-aware citations, Word/text uploads, duplicates and the size limit."""

from pathlib import Path

from docx import Document

from app.config import settings
from app.core.document_processor import DocumentProcessor
from app.core.rag_chain import page_label

SAMPLE = Path(__file__).parent / "fixtures" / "sample-leave-policy.pdf"


def upload(client, name, data, mime="application/octet-stream"):
    return client.post("/api/upload", files={"file": (name, data, mime)})


# ── chunking remembers pages ───────────────────────────────────────────


def test_chunks_record_the_pages_they_span():
    processor = DocumentProcessor(chunk_size=3, chunk_overlap=1)
    chunks = processor.chunk_pages(["a b c", "d e f g"], "x.pdf")
    assert [(c.text, c.page_start, c.page_end) for c in chunks] == [
        ("a b c", 1, 1),
        ("c d e", 1, 2),
        ("e f g", 2, 2),
        ("g", 2, 2),
    ]


def test_formats_without_pages_have_no_page_numbers():
    chunks = DocumentProcessor(chunk_size=3, chunk_overlap=1).chunk_pages(["a b c d"], "x", False)
    assert all(c.page_start is None and c.page_end is None for c in chunks)


def test_page_labels():
    assert page_label(2, 2) == "page 2"
    assert page_label(1, 2) == "pages 1–2"
    assert page_label(None, None) is None


# ── API ────────────────────────────────────────────────────────────────


def test_pdf_answers_cite_page_numbers(client):
    client.delete("/api/reset")
    body = upload(client, "sample-leave-policy.pdf", SAMPLE.read_bytes(), "application/pdf").json()
    assert body["pages"] == 3 and body["chunks_stored"] == 3

    sources = client.get("/api/ask", params={"q": "sick leave"}).json()["sources"]
    assert all(1 <= s["page_start"] <= s["page_end"] <= 3 for s in sources)
    first = next(s for s in sources if s["chunk_index"] == 0)
    assert first["page_start"] == 1


def test_same_content_under_another_name_is_rejected(client):
    res = upload(client, "copy-of-policy.pdf", SAMPLE.read_bytes(), "application/pdf")
    assert res.status_code == 409
    assert "sample-leave-policy.pdf" in res.json()["detail"]
    assert client.get("/api/stats").json()["total_documents"] == 1


def test_word_documents_are_supported(client, tmp_path):
    doc = Document()
    doc.add_heading("Travel Policy", 1)
    doc.add_paragraph("Economy class is required for flights under six hours.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Daily meal allowance"
    table.rows[0].cells[1].text = "40 dollars"
    path = tmp_path / "travel.docx"
    doc.save(path)

    body = upload(client, "travel.docx", path.read_bytes()).json()
    assert body["chunks_stored"] == 1 and body["pages"] is None

    texts = client.get("/api/search", params={"q": "meal allowance"}).json()["results"]
    travel = next(r for r in texts if r["source"] == "travel.docx")
    assert "Daily meal allowance | 40 dollars" in travel["text"]
    assert travel["page_start"] is None


def test_text_and_markdown_files_are_supported(client):
    txt = upload(client, "notes.txt", b"Remote work is allowed two days per week.", "text/plain")
    md = upload(client, "faq.md", b"# FAQ\n\nLaptops are replaced every **three** years.")
    assert txt.status_code == 200 and md.status_code == 200
    assert client.get("/api/stats").json()["total_documents"] == 4


def test_files_over_the_size_limit_are_rejected(client, monkeypatch):
    monkeypatch.setattr(settings, "MAX_UPLOAD_MB", 0.001)  # about 1 KB
    res = upload(client, "big.txt", b"x " * 2000, "text/plain")
    assert res.status_code == 413
    assert "larger than 0.001 MB" in res.json()["detail"]
    client.delete("/api/reset")
