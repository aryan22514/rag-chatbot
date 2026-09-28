from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

# File types we can read, and how to describe them to people
SUPPORTED_TYPES = {".pdf": "PDF", ".docx": "Word", ".txt": "text", ".md": "Markdown"}


@dataclass
class DocumentChunk:
    text: str
    source: str
    chunk_index: int
    page_start: int | None = None  # PDF page the chunk starts on (1-based)
    page_end: int | None = None  # PDF page the chunk ends on


class DocumentProcessor:
    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 50):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    # ── reading ─────────────────────────────────────────────────────

    def extract(self, file_path: str) -> tuple[list[str], bool]:
        """Return (text per page, whether page numbers are meaningful).

        PDFs have real pages. Word and text files are returned as a single
        block, because they have no fixed page numbers to cite.
        """
        suffix = Path(file_path).suffix.lower()
        if suffix == ".pdf":
            with open(file_path, "rb") as f:
                reader = PdfReader(f)
                return [page.extract_text() or "" for page in reader.pages], True
        if suffix == ".docx":
            return [self._docx_text(file_path)], False
        if suffix in (".txt", ".md"):
            return [Path(file_path).read_text(encoding="utf-8", errors="replace")], False
        raise ValueError(f"Unsupported file type: {suffix}")

    def extract_pdf(self, file_path: str) -> str:
        """Whole PDF as one string (kept for scripts that only need the text)."""
        pages, _ = self.extract(file_path)
        return "\n".join(pages)

    @staticmethod
    def _docx_text(file_path: str) -> str:
        doc = Document(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    # ── chunking ────────────────────────────────────────────────────

    def chunk_pages(self, pages: list[str], source: str, paged: bool = True) -> list[DocumentChunk]:
        """Chunk across page boundaries, remembering which pages each chunk covers."""
        words: list[str] = []
        page_of: list[int] = []
        for number, page in enumerate(pages, start=1):
            page_words = page.split()
            words.extend(page_words)
            page_of.extend([number] * len(page_words))
        return self._chunk(words, source, page_of if paged else None)

    def chunk_text(self, text: str, source: str) -> list[DocumentChunk]:
        return self._chunk(text.split(), source, None)

    def _chunk(
        self, words: list[str], source: str, page_of: list[int] | None
    ) -> list[DocumentChunk]:
        if not words:
            return []

        chunks = []
        stride = self.chunk_size - self.chunk_overlap

        for i, start in enumerate(range(0, len(words), stride)):
            end = min(start + self.chunk_size, len(words))
            piece = " ".join(words[start:end]).strip()

            if piece:
                chunks.append(
                    DocumentChunk(
                        text=piece,
                        source=source,
                        chunk_index=i,
                        page_start=page_of[start] if page_of else None,
                        page_end=page_of[end - 1] if page_of else None,
                    )
                )

        return chunks
